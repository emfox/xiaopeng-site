#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
export_amap.py —— 把小鹏站点导出成「高德地图小程序」可批量导入的 Excel

高德 App → 我的 → 地图小程序 → 在 PC 端 (wia.amap.com) 批量导入数据，
导入后点位就在高德 App 本体里，支持一键导航、顺路规划。

模板列（按高德实际模板，带 * 的是必填）：
    名称 | *经度 | *纬度 | *地址 | 颜色 | 图标(外轮廓) | 图标(填充物) | 描述 | 文件夹

注意：经纬度必须是「经度」「纬度」两列，合成一列 "lng,lat" 高德会报
「不支持的文件格式」。上过这个当，别再改回去。

颜色（前端色块顺序，1 开始）：
    1 绿色  免费权益 + 有停车优惠   ← 最想要的一类
    2 蓝色  免费权益、无停车优惠
    3 橙色  第三方合作 + 有停车优惠
    4 灰色  第三方合作、无停车优惠

用法：
  python3 export_amap.py                 # 每个城市一个 xlsx → amap/
  python3 export_amap.py --all           # 额外生成全部城市合一的文件
  python3 export_amap.py --cities 上海,杭州
  python3 export_amap.py --park-only     # 只要带停车优惠的
"""
from __future__ import annotations

import argparse
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "amap")

# 高德模板表头。高德界面上显示的是「*经度」「*纬度」「*地址」，星号只是必填标记，
# 实际模板文件里通常不带星号。用不带星号的写法兼容性更好。
HEADER = ["名称", "经度", "纬度", "地址", "颜色",
          "图标(外轮廓)", "图标(填充物)", "描述", "文件夹"]
WIDTHS = [42, 12, 12, 52, 6, 14, 14, 46, 10]

COLOR_FREE_PARK = 1
COLOR_FREE = 2
COLOR_THIRD_PARK = 3
COLOR_THIRD = 4

COLOR_NAME = {
    COLOR_FREE_PARK: "绿",
    COLOR_FREE: "蓝",
    COLOR_THIRD_PARK: "橙",
    COLOR_THIRD: "灰",
}


def color_of(p) -> int:
    free = p.get("type") != "third"
    park = bool(p.get("park"))
    if free and park:
        return COLOR_FREE_PARK
    if free:
        return COLOR_FREE
    if park:
        return COLOR_THIRD_PARK
    return COLOR_THIRD


def clean_price(price) -> str:
    """价格形如 "0.64+0.61"（电费+服务费）。接口有时返回 "None+None"，要清掉。"""
    if not price:
        return ""
    s = str(price).replace("None", "").replace("none", "")
    parts = [x.strip() for x in s.split("+") if x.strip()]
    if not parts:
        return ""
    if len(parts) == 1:
        return parts[0]
    return f"电费{parts[0]} 服务费{parts[1]}"


def describe(p) -> str:
    bits = []
    if p.get("type") == "third":
        bits.append(f"第三方·{p.get('oper') or '合作站'}")
    else:
        bits.append("小鹏免费权益")
    if p.get("park"):
        bits.append("有停车优惠")
    if p.get("guns"):
        bits.append(p["guns"])
    price = clean_price(p.get("price"))
    if price:
        bits.append(price)
    return "；".join(bits)


def rows_for(points):
    rows = []
    for p in points:
        lng, lat = p.get("lng"), p.get("lat")
        if lng is None or lat is None:
            continue
        rows.append([
            p.get("name", ""),
            round(float(lng), 6),
            round(float(lat), 6),
            p.get("addr", ""),
            color_of(p),
            "",            # 图标(外轮廓)：留空用默认
            "",            # 图标(填充物)：留空用默认
            describe(p),
            "",            # 文件夹：留空，不分组
        ])
    return rows


def write_xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")
    for c, name in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, r in enumerate(rows, 2):
        for c, v in enumerate(r, 1):
            ws.cell(row=i, column=c, value=v)
        # 经纬度强制数字格式，避免被 Excel 当文本
        for c in (2, 3):
            ws.cell(row=i, column=c).number_format = "0.000000"

    for c, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
    return len(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cities", default="", help="只导出这些城市，逗号分隔")
    ap.add_argument("--all", action="store_true", help="额外生成全部城市合一的大文件")
    ap.add_argument("--park-only", action="store_true", help="只导出有停车优惠的站")
    ap.add_argument("--out", default=OUT, help="输出目录，默认 amap/")
    ap.add_argument("--csv", action="store_true", help="同时另存一份 csv（备用）")
    a = ap.parse_args()

    src = os.path.join(ROOT, "data.json")
    if not os.path.exists(src):
        sys.exit("找不到 data.json，先跑 python3 build.py")
    data = json.load(open(src, encoding="utf-8"))
    points = data["points"]

    if a.park_only:
        points = [p for p in points if p.get("park")]
    if a.cities:
        want = {c.strip() for c in a.cities.split(",") if c.strip()}
        points = [p for p in points if p.get("city") in want]

    by_city = {}
    for p in points:
        by_city.setdefault(p.get("city") or "未知", []).append(p)

    os.makedirs(a.out, exist_ok=True)
    print(f"导出目录：{a.out}")

    for old in os.listdir(a.out):
        if old.endswith((".csv", ".xlsx")):
            os.remove(os.path.join(a.out, old))

    total = 0
    for city in sorted(by_city, key=lambda c: -len(by_city[c])):
        rows = rows_for(by_city[city])
        n = write_xlsx(os.path.join(a.out, f"{city}.xlsx"), rows)
        total += n
        print(f"  {city}.xlsx  {n} 个站")

    if a.all:
        n = write_xlsx(os.path.join(a.out, "全部城市.xlsx"), rows_for(points))
        print(f"  全部城市.xlsx  {n} 个站")

    if a.csv:
        import csv
        for city, arr in by_city.items():
            with open(os.path.join(a.out, f"{city}.csv"), "w",
                      encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                w.writerow(HEADER)
                w.writerows(rows_for(arr))
        print("  另已生成 csv 备份")

    with open(os.path.join(a.out, "导入说明.txt"), "w", encoding="utf-8") as f:
        f.write(
            "导入方法\n"
            "1. 电脑浏览器打开 https://wia.amap.com 并登录高德账号\n"
            "2. 左上「我加入的」→ 创建新地图，起个名字（如「小鹏免费站」）\n"
            "3. 点顶部「批量导入」→ 数据来源选「高德地图」→ 上传这里的 .xlsx\n"
            "4. 手机高德 App → 我的 → 地图小程序 → 就能看到，可一键导航\n\n"
            "为什么必须用 xlsx\n"
            "  高德要求经纬度是「经度」「纬度」两列。合成一列 \"lng,lat\" 会报\n"
            "  「不支持的文件格式」，是这个错误最常见的成因。\n\n"
            "列格式（与高德模板一致）\n"
            "  名称  *经度  *纬度  *地址  颜色  图标(外轮廓)  图标(填充物)  描述  文件夹\n\n"
            "颜色含义\n"
            "  1 绿色   小鹏免费权益 + 有停车优惠\n"
            "  2 蓝色   小鹏免费权益、无停车优惠\n"
            "  3 橙色   第三方合作 + 有停车优惠\n"
            "  4 灰色   第三方合作、无停车优惠\n\n"
            "注意\n"
            "  - 单次上传 ≤ 50000 行、≤ 20 MB，单城市最多 224 站，远低于上限\n"
            "  - 高德目前不支持覆盖更新，重新导入前要先删旧地图或新建一张\n"
            "  - 坐标系 GCJ-02，和高德一致，不用转换\n"
        )
    print("  导入说明.txt")
    print(f"\n合计 {len(by_city)} 个城市 · {total} 个站")


if __name__ == "__main__":
    main()
